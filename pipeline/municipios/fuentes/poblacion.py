"""Población (DANE): proyecciones municipales por área geográfica, 2018-2042.

El DANE publica un Excel con una fila por municipio, año y área (cabecera, centros poblados y
rural disperso, total). Versión vigente: actualizada el 30 de julio de 2025 (post-COVID).
"""

from collections.abc import Iterable, Iterator
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from scrapling.fetchers import Fetcher

from municipios.comun.modelos import Observacion

FUENTE = "dane_proyecciones_poblacion"
URL = (
    "https://www.dane.gov.co/files/censo2018/proyecciones-de-poblacion/"
    "Municipal/PPED-AreaMun-2018-2042_VP.xlsx"
)
HOJA = "PobMunicipalxÁrea"

# Área geográfica del DANE -> indicador del proyecto.
INDICADORES = {
    "Total": "poblacion",
    "Cabecera Municipal": "poblacion_cabecera",
    "Centros Poblados y Rural Disperso": "poblacion_rural",
}


def descargar(directorio: Path, refrescar: bool = False) -> Path:
    """Descarga el Excel (unos 4 MB) a la caché y devuelve su ruta."""
    ruta = directorio / Path(URL).name
    if ruta.exists() and not refrescar:
        return ruta
    respuesta = Fetcher.get(URL, timeout=300, retries=3)
    if respuesta.status != 200:
        raise RuntimeError(f"HTTP {respuesta.status} al descargar {URL}")
    ruta.parent.mkdir(parents=True, exist_ok=True)
    ruta.write_bytes(respuesta.body)
    return ruta


def leer_filas(ruta: Path) -> Iterator[dict[str, Any]]:
    """Lee la hoja de datos y produce {codigo, anio, area, total} por fila.

    Ubica las columnas por el encabezado. Ojo: en el Excel la columna «MPIO» trae el código de
    5 dígitos y «DPMP» trae el nombre (los rótulos están cruzados), así que se usa «MPIO».
    """
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        filas = libro[HOJA].iter_rows(values_only=True)
        columnas: dict[str, int] | None = None
        for fila in filas:
            if columnas is None:
                if fila and "AÑO" in fila:
                    columnas = {nombre: i for i, nombre in enumerate(fila) if nombre}
                continue
            codigo = fila[columnas["MPIO"]]
            if not (isinstance(codigo, str) and len(codigo) == 5 and codigo.isdigit()):
                continue  # filas vacías y notas al pie
            yield {
                "codigo": codigo,
                "anio": int(fila[columnas["AÑO"]]),
                "area": str(fila[columnas["ÁREA GEOGRÁFICA"]]).strip(),
                "total": int(fila[columnas["TOTAL"]]),
            }
        if columnas is None:
            raise ValueError(f"No se encontró el encabezado en la hoja {HOJA!r}")
    finally:
        libro.close()


def normalizar(
    filas: Iterable[dict[str, Any]], divipolas: set[str]
) -> tuple[list[Observacion], dict[str, int]]:
    """Observaciones de población por municipio, año y área.

    Devuelve también los códigos que no cruzan con DIVIPOLA (con su población total más
    reciente), para revisarlos en vez de descartarlos en silencio.
    """
    observaciones: list[Observacion] = []
    sin_cruce: dict[str, int] = {}
    for fila in filas:
        indicador = INDICADORES.get(fila["area"])
        if indicador is None:
            raise ValueError(f"Área geográfica desconocida: {fila['area']!r}")
        if fila["codigo"] not in divipolas:
            if indicador == "poblacion":
                sin_cruce[fila["codigo"]] = fila["total"]
            continue
        observaciones.append(
            Observacion(
                divipola=fila["codigo"],
                indicador=indicador,
                valor=fila["total"],
                periodo=str(fila["anio"]),
                fuente=FUENTE,
            )
        )
    observaciones.sort(key=lambda o: (o.divipola, o.indicador, o.periodo))
    return observaciones, sin_cruce
